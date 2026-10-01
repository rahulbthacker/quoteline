"""Exports: the comparison as XLSX and CSV, an analysis result, and the award pack as PDF."""

from __future__ import annotations

import csv
import io
from datetime import UTC, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.grid import GridResult, money

FILLS = {
    "needs_review": PatternFill("solid", fgColor="FBF0D3"),
    "not_quoted": PatternFill("solid", fgColor="EDEAE2"),
    "overridden": PatternFill("solid", fgColor="FBEAE2"),
    "lowest": PatternFill("solid", fgColor="E3F0E4"),
}
HEAD = PatternFill("solid", fgColor="F0EDE6")


def _cell_text(cell) -> str | float:
    if cell.excluded:
        return "excluded"
    if cell.state == "not_quoted":
        return "not quoted"
    if cell.state == "needs_review":
        return f"in review ({money(cell.display_inr)}?)" if cell.display_inr else "in review"
    return cell.inr_per_base


def grid_xlsx(result: GridResult) -> bytes:
    grid = result.grid
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Comparison"
    a = grid.assumptions
    sheet.append([f"{grid.rfx_id}: {grid.title}"])
    sheet.append([f"Normalised to {a.base_unit}. FX: {a.fx}. {a.freight} {a.tax}"])
    sheet.append(["Cells in review, not quoted or excluded are NOT in any total. "
                  f"Exported {datetime.now(UTC).strftime('%d %b %Y %H:%M UTC')}."])  # fmt: skip
    sheet.append([])
    header = ["Line", "Description", "Qty", "UOM"] + [v.name for v in grid.vendors] + ["Last year"]
    sheet.append(header)
    for cell in sheet[5]:
        cell.font, cell.fill = Font(bold=True), HEAD
    for row in grid.lines:
        values = [row.line_id, row.description, row.qty, row.uom]
        values += [_cell_text(row.cells[v.vendor_id]) for v in grid.vendors] + [row.last_year_inr]
        sheet.append(values)
        for index, vendor in enumerate(grid.vendors):
            cell = row.cells[vendor.vendor_id]
            target = sheet.cell(sheet.max_row, 5 + index)
            target.number_format = "#,##0"
            key = "lowest" if cell.is_lowest else cell.state
            if key in FILLS:
                target.fill = FILLS[key]
            if not isinstance(target.value, int | float):
                target.alignment = Alignment(horizontal="right")
    sheet.append([])
    sheet.append(
        ["", "Extended total, counted lines only (₹)", "", ""]
        + [v.totals.extended_inr for v in grid.vendors]
    )
    sheet.append(["", "Lines counted", "", ""] + [v.totals.lines_counted for v in grid.vendors])
    sheet.append(
        ["", "Lines in review (excluded)", "", ""] + [v.totals.needs_review for v in grid.vendors]
    )
    sheet.append(
        ["", "Lines not quoted (excluded)", "", ""] + [v.totals.not_quoted for v in grid.vendors]
    )
    sheet.append(
        ["", "Questionnaire", "", ""] + [v.quality or "not evaluated" for v in grid.vendors]
    )
    sheet.append(["", "Freight, as stated", "", ""] + [
        "included" if v.freight.kind == "included" else v.freight.amount_inr if v.freight.amount_inr
        else "extra, unpriced" if v.freight.kind == "extra_unpriced" else "not stated"
        for v in grid.vendors
    ])  # fmt: skip
    for r in range(sheet.max_row - 5, sheet.max_row + 1):
        sheet.cell(r, 2).font = Font(bold=True)
        for c in range(5, 5 + len(grid.vendors)):
            sheet.cell(r, c).number_format = "#,##0"
    sheet.column_dimensions["B"].width = 62
    for index in range(5, 6 + len(grid.vendors)):
        sheet.column_dimensions[get_column_letter(index)].width = 24
    sheet.freeze_panes = "E6"

    provenance = workbook.create_sheet("Provenance")
    provenance.append(["Line", "Vendor", "State", "₹ per 1,000 (counted)", "As quoted", "Source file",
                       "Locator", "Verbatim snippet", "Model", "Confidence", "Agreement", "Arithmetic",
                       "Flag / override reason"])  # fmt: skip
    for cell in provenance[1]:
        cell.font, cell.fill = Font(bold=True), HEAD
    names = {v.vendor_id: v.name for v in grid.vendors}
    for row in grid.lines:
        for vendor_id, cell in row.cells.items():
            steps = " → ".join(
                f"{s.op} ×{s.factor:g} = {s.value_out:,.2f} [{s.source}]"
                for s in cell.normalization_trace
            )
            reason = cell.override.reason if cell.override else cell.flag_reason
            readings = cell.readings or [None]
            for reading in readings:
                provenance.append([
                    row.line_id, names[vendor_id], "excluded" if cell.excluded else cell.state,
                    cell.inr_per_base, cell.raw_display,
                    reading.source_ref.file if reading else None,
                    reading.source_ref.locator if reading else None,
                    reading.source_ref.snippet if reading else None,
                    reading.model if reading else None,
                    reading.confidence if reading else None, cell.agreement, steps, reason,
                ])  # fmt: skip
    for column, width in zip(
        "ABCDEFGHIJKLM", (7, 24, 14, 18, 22, 34, 28, 70, 26, 11, 10, 70, 60), strict=True
    ):
        provenance.column_dimensions[column].width = width

    review = workbook.create_sheet("Review items")
    review.append(["Vendor", "Kind", "Lines", "Status", "Reason", "Resolution"])
    for cell in review[1]:
        cell.font, cell.fill = Font(bold=True), HEAD
    for item in result.items:
        resolution = (
            (item.resolution or {}).get("label") or (item.resolution or {}).get("option_id") or ""
        )
        review.append([item.vendor_name, item.title, ", ".join(r.line_id for r in item.cells),
                       item.status, item.reason, resolution])  # fmt: skip
    for column, width in zip("ABCDEF", (24, 34, 30, 14, 100, 50), strict=True):
        review.column_dimensions[column].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def grid_csv(result: GridResult) -> bytes:
    """Long format: one row per cell, with state and source, for anyone's own analysis."""
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["rfx_id", "line_id", "description", "qty", "uom", "vendor", "state", "counted",
                     "inr_per_1000", "extended_inr", "as_quoted", "source_file", "locator", "reason"])  # fmt: skip
    names = {v.vendor_id: v.name for v in result.grid.vendors}
    for row in result.grid.lines:
        for vendor_id, cell in row.cells.items():
            source = cell.readings[0].source_ref if cell.readings else None
            writer.writerow([
                result.grid.rfx_id, row.line_id, row.description, row.qty, row.uom, names[vendor_id],
                "excluded" if cell.excluded else cell.state, cell.inr_per_base is not None,
                cell.inr_per_base, cell.extended_inr, cell.raw_display,
                source.file if source else "", source.locator if source else "",
                (cell.override.reason if cell.override else cell.flag_reason) or "",
            ])  # fmt: skip
    return out.getvalue().encode("utf-8-sig")


def rows_xlsx(
    title: str, question: str, sql: str | None, columns: list[str], rows: list[list], note: str
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Result"
    sheet.append([title])
    sheet.append([f"Question: {question}"])
    sheet.append([note])
    sheet.append([])
    sheet.append(columns)
    for cell in sheet[5]:
        cell.font, cell.fill = Font(bold=True), HEAD
    for row in rows:
        sheet.append(row)
    for index in range(1, len(columns) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 22
    query = workbook.create_sheet("Query")
    query.append(["The query that produced this result"])
    for line in (sql or "").splitlines():
        query.append([line])
    query.column_dimensions["A"].width = 140
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def rows_csv(columns: list[str], rows: list[list]) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(columns)
    writer.writerows(rows)
    return out.getvalue().encode("utf-8-sig")


def award_pdf(
    result: GridResult, scenario_name: str, award: dict, audit_log: list[dict], status: str
) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    def rupees(value: float | None) -> str:
        return "-" if value is None else money(value).replace("₹", "Rs ")

    h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=15, leading=19)
    h2 = ParagraphStyle(
        "h2", fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=9, spaceAfter=3
    )
    body = ParagraphStyle("body", fontName="Helvetica", fontSize=8.8, leading=12)
    small = ParagraphStyle(
        "small",
        fontName="Helvetica",
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor("#555555"),
    )
    grid = result.grid
    a = grid.assumptions

    def clean(text: str) -> str:
        return text.replace("₹", "Rs ").replace("&", "&amp;").replace("<", "&lt;")

    story = [
        Paragraph(f"Award recommendation: {grid.rfx_id}", h1),
        Paragraph(
            clean(
                f"{grid.title}. Scenario: {scenario_name}. Status: {status}. "
                f"Generated {datetime.now(UTC).strftime('%d %b %Y %H:%M UTC')}."
            ),
            body,
        ),  # fmt: skip
        Spacer(1, 3 * mm),
        Paragraph(
            clean(
                f"Recommended total {rupees(award['total_inr'])} across {award['lines_allocated']} lines and "
                f"{award['vendors_awarded']} vendors."
            ),
            body,
        ),  # fmt: skip
    ]
    if award["best_single_vendor"]:
        best = award["best_single_vendor"]
        story.append(Paragraph(clean(
            f"Against the best single vendor covering the same lines ({best['vendor']}, {rupees(best['extended_inr'])}): "
            f"{rupees(award['vs_best_single_inr'])} ({award['vs_best_single_pct']}%)."), body))  # fmt: skip
    ly = award["last_year"]
    if ly["delta_pct"] is not None:
        story.append(Paragraph(clean(
            f"Against last year's contract on {ly['lines_compared']} like-for-like lines: "
            f"{rupees(ly['delta_inr'])} ({ly['delta_pct']}%)."), body))  # fmt: skip
    if award["open_items"]:
        story.append(Paragraph(clean(
            f"NOT READY FOR APPROVAL: {len(award['open_items'])} review item(s) are open. "
            f"Cells in review: {', '.join(award['cells_in_review']) or 'none'}."), body))  # fmt: skip

    story.append(Paragraph("Allocation", h2))
    rows = [["Vendor", "Quality", "Lines", "Extended (Rs)", "Share", "Freight"]]
    for item in award["allocations"]:
        freight = item["freight"]
        freight_text = ("included" if freight["kind"] == "included" else rupees(freight["amount_inr"])
                        if freight["amount_inr"] else "unpriced")  # fmt: skip
        rows.append([item["vendor"], item["quality"] or "-", str(item["line_count"]),
                     rupees(item["extended_inr"]), f"{item['share_pct']}%", freight_text])  # fmt: skip
    table_style = TableStyle([
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.5), ("FONT", (0, 1), (-1, -1), "Helvetica", 8.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F0EDE6")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BDB6A6")), ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
    ])  # fmt: skip
    story.append(
        Table(
            rows,
            colWidths=[55 * mm, 20 * mm, 14 * mm, 34 * mm, 18 * mm, 30 * mm],
            style=table_style,
        )
    )

    story.append(Paragraph("Rules applied", h2))
    story += [Paragraph(clean(f"- {rule}"), body) for rule in award["rules_applied"]]
    story.append(Paragraph("Assumptions in force", h2))
    story.append(
        Paragraph(clean(f"Base unit: {a.base_unit}. FX: {a.fx}. {a.freight} {a.tax}"), body)
    )
    if award["exclusions"]:
        story.append(Paragraph("Excluded and why", h2))
        story += [
            Paragraph(clean(f"- {e['vendor']}: {e['reason']}"), body) for e in award["exclusions"]
        ]
        if award["rule_cost_text"]:
            story.append(Paragraph(clean(award["rule_cost_text"]), body))
    if award["unallocated"]:
        story.append(Paragraph("Lines not allocated", h2))
        story += [
            Paragraph(clean(f"- {u['line_id']} {u['description']}: {u['reason']}"), body)
            for u in award["unallocated"]
        ]

    story.append(Paragraph("Line allocation", h2))
    line_rows = [["Line", "Vendor", "Rs per 1,000", "Extended (Rs)"]]
    for item in award["line_allocations"]:
        line_rows.append(
            [
                item["line_id"],
                item["vendor"],
                rupees(item["inr_per_1000"]),
                rupees(item["extended_inr"]),
            ]
        )
    story.append(
        Table(
            line_rows,
            colWidths=[16 * mm, 70 * mm, 40 * mm, 45 * mm],
            style=table_style,
            repeatRows=1,
        )
    )

    story.append(Paragraph("Audit log", h2))
    for event in audit_log[:40]:
        when = event["at"][:16].replace("T", " ")
        story.append(Paragraph(clean(f"{when} UTC · {event['actor']} · {event['summary']}"), small))
    document = SimpleDocTemplate(io.BytesIO(), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                                 topMargin=16 * mm, bottomMargin=16 * mm, title=f"Award {grid.rfx_id}")  # fmt: skip
    buffer = io.BytesIO()
    document.filename = buffer
    document.build(story)
    return buffer.getvalue()
