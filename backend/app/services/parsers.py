"""File parsers. They produce text with locators plus page images; the model does the reading.

Nothing here interprets a price. A parser's only job is to hand the model the content of a
file in a form where every value has an address (sheet and cell, page, paragraph, email line)
that the model can cite and a human can check.
"""

from __future__ import annotations

import email
import email.policy
import io
import zipfile
from dataclasses import dataclass, field

import pypdf
import pypdfium2 as pdfium
from docx import Document
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image

from app.llm.contracts import ImagePart
from app.schemas.reply import FileFormat

MAX_IMAGE_EDGE = 2576  # the largest edge the vision models use at full resolution
PDF_RENDER_SCALE = 2.0  # 144 dpi
PREVIEW_CHARS = 1800


class UnsupportedFile(Exception):
    pass


@dataclass
class ParsedDocument:
    name: str
    format: FileFormat
    text: str = ""
    images: list[ImagePart] = field(default_factory=list)
    image_locators: list[str] = field(default_factory=list)
    sections: list[str] = field(default_factory=list)
    pages: int | None = None
    children: list[tuple[str, bytes]] = field(default_factory=list)  # files carried inside
    cells: dict[str, object] = field(default_factory=dict)  # "Sheet!A1" -> value (workbooks)

    @property
    def preview(self) -> str:
        return self.text[:PREVIEW_CHARS]


def sniff_format(name: str, data: bytes) -> FileFormat:
    """Decide the format from the bytes, falling back to the extension."""
    if data[:4] == b"%PDF":
        return "pdf"
    if data[:3] == b"\xff\xd8\xff" or data[:8] == b"\x89PNG\r\n\x1a\n":
        return "image"
    if data[:2] == b"PK":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = archive.namelist()
        if any(n.startswith("xl/") for n in names):
            return "xlsx"
        if any(n.startswith("word/") for n in names):
            return "docx"
    lowered = name.lower()
    if lowered.endswith(".eml") or data[:5].lower() in (b"from:", b"recei", b"retur", b"mime-"):
        return "eml"
    raise UnsupportedFile(f"{name}: unrecognised file type")


def parse(name: str, data: bytes) -> ParsedDocument:
    fmt = sniff_format(name, data)
    return _PARSERS[fmt](name, data)


def _cell_text(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\n", " ")


def _parse_xlsx(name: str, data: bytes) -> ParsedDocument:
    workbook = load_workbook(io.BytesIO(data), data_only=True)
    parts: list[str] = []
    cell_map: dict[str, object] = {}
    for sheet in workbook.worksheets:
        parts.append(
            f'=== Sheet "{sheet.title}" (rows 1-{sheet.max_row}, '
            f"columns A-{get_column_letter(sheet.max_column)}) ==="
        )
        merged = sorted(str(rng) for rng in sheet.merged_cells.ranges)
        parts.append(
            f"Cite cells on this sheet as {sheet.title}!<cell>, for example {sheet.title}!A1."
        )
        if merged:
            parts.append("Merged ranges: " + ", ".join(merged))
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cell_map[f"{sheet.title}!{cell.coordinate}"] = cell.value
            cells = [
                f'{cell.coordinate}="{_cell_text(cell.value)}"'
                for cell in row
                if cell.value is not None and str(cell.value).strip()
            ]
            if cells:
                parts.append(" | ".join(cells))
        parts.append("")
    return ParsedDocument(
        name=name,
        format="xlsx",
        text="\n".join(parts),
        sections=[sheet.title for sheet in workbook.worksheets],
        cells=cell_map,
    )


def _parse_docx(name: str, data: bytes) -> ParsedDocument:
    document = Document(io.BytesIO(data))
    parts: list[str] = []
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            if paragraph.text.strip():
                parts.append(f"[header] {paragraph.text.strip()}")
    number = 0
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            number += 1
            parts.append(f"[paragraph {number}] {paragraph.text.strip()}")
    for t_index, table in enumerate(document.tables, start=1):
        for r_index, row in enumerate(table.rows, start=1):
            cells = " | ".join(cell.text.strip() for cell in row.cells)
            parts.append(f"[table {t_index} row {r_index}] {cells}")
    return ParsedDocument(name=name, format="docx", text="\n".join(parts), sections=["body"])


def _fit(image: Image.Image) -> Image.Image:
    longest = max(image.size)
    if longest <= MAX_IMAGE_EDGE:
        return image
    ratio = MAX_IMAGE_EDGE / longest
    return image.resize((round(image.width * ratio), round(image.height * ratio)), Image.LANCZOS)


def _jpeg(image: Image.Image, quality: int = 90) -> ImagePart:
    buffer = io.BytesIO()
    image.convert("RGB").save(buffer, "JPEG", quality=quality)
    return ImagePart(media_type="image/jpeg", data=buffer.getvalue())


def _parse_pdf(name: str, data: bytes) -> ParsedDocument:
    reader = pypdf.PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        parts.append(f"=== page {index} ===")
        parts.append((page.extract_text() or "").strip())
        parts.append("")
    rendered = pdfium.PdfDocument(data)
    images = [_jpeg(_fit(page.render(scale=PDF_RENDER_SCALE).to_pil())) for page in rendered]
    count = len(reader.pages)
    return ParsedDocument(
        name=name,
        format="pdf",
        text="\n".join(parts),
        images=images,
        image_locators=[f"page {n}" for n in range(1, count + 1)],
        sections=[f"page {n}" for n in range(1, count + 1)],
        pages=count,
    )


def _parse_image(name: str, data: bytes) -> ParsedDocument:
    image = Image.open(io.BytesIO(data))
    image.load()
    return ParsedDocument(
        name=name,
        format="image",
        text="",
        images=[_jpeg(_fit(image), quality=92)],
        image_locators=["image"],
        sections=["image"],
        pages=1,
    )


def _parse_eml(name: str, data: bytes) -> ParsedDocument:
    message = email.message_from_bytes(data, policy=email.policy.default)
    parts = [f"[email header] {key}: {message[key]}" for key in ("From", "To", "Subject", "Date")]
    body = message.get_body(preferencelist=("plain", "html"))
    lines = [line.strip() for line in (body.get_content() if body else "").splitlines()]
    for number, line in enumerate((line for line in lines if line), start=1):
        parts.append(f"[email body, line {number}] {line}")
    children = [
        (part.get_filename() or "attachment", part.get_content())
        for part in message.iter_attachments()
        if isinstance(part.get_content(), bytes)
    ]
    for child_name, _payload in children:
        parts.append(f"[email attachment] {child_name}")
    return ParsedDocument(
        name=name, format="eml", text="\n".join(parts), sections=["email body"], children=children
    )


_PARSERS = {
    "xlsx": _parse_xlsx,
    "docx": _parse_docx,
    "pdf": _parse_pdf,
    "image": _parse_image,
    "eml": _parse_eml,
}
