"""Checkpoint B: the golden set matches the generators and the files on disk.

These tests read the committed reply files with plain parsers (no model involved) and check
that every golden value is physically present where the golden set says it is. They are the
ground the extraction tests in Checkpoint C stand on.
"""

from __future__ import annotations

import email
import email.policy
import hashlib
import json
import re
import sys
from pathlib import Path

import pypdf
import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image

from app.schemas import RFx

REPO = Path(__file__).resolve().parents[2]
DATA = REPO / "data"
sys.path.insert(0, str(DATA / "gen"))

import catalog  # noqa: E402

VENDORS = ["meridian", "sribalaji", "kaveri", "northbox", "anand"]
FX = 83.42


def golden(vendor: str) -> dict:
    return json.loads((DATA / "golden" / f"{vendor}.json").read_text(encoding="utf-8"))


def reply_path(vendor: str) -> Path:
    return DATA / "replies" / vendor / golden(vendor)["files"][0]


@pytest.fixture(scope="module")
def rfx() -> RFx:
    return RFx.model_validate_json((DATA / "rfx" / "RFX-0147.json").read_text(encoding="utf-8"))


def test_rfx_meets_demo_scale(rfx: RFx):
    assert [line.line_id for line in rfx.line_items] == [f"L{n:02d}" for n in range(1, 31)]
    assert len(rfx.questionnaire) == 8
    assert sum(q.mandatory for q in rfx.questionnaire) == 3
    assert rfx.fx_policy.rate == FX
    new_lines = [line.line_id for line in rfx.line_items if line.last_year_price_inr is None]
    assert new_lines == ["L25", "L27"]
    assert all(line.spec.board_weight_kg_per_pc for line in rfx.line_items)


def test_contract_value_and_price_range_are_credible(rfx: RFx):
    """About ₹4 crore, with rates from roughly ₹1.6k (pads) to ₹1.2 L (7-ply export cartons)."""
    qty = {line.line_id: line.qty for line in rfx.line_items}
    for vendor in ("meridian", "sribalaji"):
        lines = golden(vendor)["lines"]
        total = sum(row["normalised_inr_per_1000"] * qty[row["line_id"]] / 1000 for row in lines)
        assert 3.5e7 < total < 4.0e7
        rates = [row["normalised_inr_per_1000"] for row in lines]
        assert 1_500 < min(rates) and max(rates) < 125_000


@pytest.mark.parametrize(
    ("vendor", "quoted"),
    [("meridian", 30), ("sribalaji", 30), ("kaveri", 29), ("northbox", 27), ("anand", 30)],
)
def test_golden_shape(vendor: str, quoted: int):
    data = golden(vendor)
    assert [row["line_id"] for row in data["lines"]] == [f"L{n:02d}" for n in range(1, 31)]
    assert data["lines_quoted_count"] == quoted == sum(row["quoted"] for row in data["lines"])
    for row in data["lines"]:
        if row["expected_state"] == "not_quoted":
            assert row["normalised_inr_per_1000"] is None
        if row["quoted"]:
            assert row["locator"] and row["raw_text"]


def test_golden_matches_generator():
    """The golden JSON is exactly what the deterministic catalog produces."""
    quotes = catalog.build_quotes()
    for vendor in VENDORS:
        for row in golden(vendor)["lines"]:
            quote = quotes[vendor][row["line_id"]]
            assert row["raw_value"] == quote.raw_value
            assert row["normalised_inr_per_1000"] == quote.true_inr_per_1000


def test_manifest_hashes_match_files():
    manifest = json.loads((DATA / "replies" / "manifest.json").read_text())
    assert [entry["vendor_id"] for entry in manifest] == VENDORS
    for entry in manifest:
        for item in entry["files"] + entry["attachments"]:
            assert hashlib.sha256((REPO / item["path"]).read_bytes()).hexdigest() == item["sha256"]
    loose = [Path(a["path"]).name for entry in manifest for a in entry["attachments"]]
    assert sorted(loose) == ["Meridian_Spec_Sheet.pdf", "SriBalaji_Ink_MSDS.pdf"]


def test_meridian_xlsx_holds_every_rate_at_its_cell():
    data = golden("meridian")
    workbook = load_workbook(reply_path("meridian"))
    assert workbook.sheetnames[1] == "Commercial Offer"  # prices are on the second sheet
    sheet = workbook["Commercial Offer"]
    assert any(str(rng) == "D5:F5" for rng in sheet.merged_cells.ranges)  # merged header
    assert sheet["I5"].value == "Rate"  # no unit in the header
    note_sheet, note_cell = data["unit_note"]["locator"].split("!")
    assert "INR per piece" in workbook[note_sheet][note_cell].value
    for row in data["lines"]:
        sheet_name, cell = row["locator"].split("!")
        assert workbook[sheet_name][cell].value == row["raw_value"]
        assert row["normalised_inr_per_1000"] == pytest.approx(row["raw_value"] * 1000)
    text = " ".join(str(c.value) for r in sheet.iter_rows() for c in r if c.value is not None)
    assert not re.search(r"\bL\d\d\b", text)  # ignores the template: no RFx line ids


def test_sribalaji_pdf_has_three_pages_rates_and_buried_footnote():
    data = golden("sribalaji")
    reader = pypdf.PdfReader(reply_path("sribalaji"))
    assert len(reader.pages) == 3
    pages = [page.extract_text() for page in reader.pages]
    for row in data["lines"]:
        page = int(re.match(r"page (\d)", row["locator"]).group(1))
        assert row["raw_text"] in pages[page - 1]
        assert float(row["raw_text"].replace(",", "")) == row["raw_value"]
    assert "2. 4% discount on lines above ₹1 L." in pages[1]
    assert data["discount"]["qualifying_lines_by_rate"] == ["L04"]
    l04 = next(row for row in data["lines"] if row["line_id"] == "L04")
    assert l04["expected_state"] == "needs_review"
    assert l04["true_inr_per_1000"] == pytest.approx(l04["normalised_inr_per_1000"] * 0.96)
    assert "FSSC 22000" in pages[2]


def test_kaveri_docx_is_prose_in_usd_with_the_per_box_sentence():
    data = golden("kaveri")
    document = Document(reply_path("kaveri"))
    assert len(document.tables) == 0  # commercials are prose, not a table
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    for row in data["lines"]:
        number = int(row["locator"].split()[1])
        assert row["raw_text"] in paragraphs[number - 1]
        if row["quoted"]:
            assert row["raw_currency"] == "USD"
            factor = 1000 if row["raw_unit"] == "per_box" else 1
            assert row["normalised_inr_per_1000"] == pytest.approx(
                row["raw_value"] * factor * FX, abs=0.01
            )
    l01 = next(row for row in data["lines"] if row["line_id"] == "L01")
    sentence = paragraphs[int(l01["locator"].split()[1]) - 1]
    assert "per box, packed 100 per bundle" in sentence
    assert l01["expected_state"] == "needs_review"
    assert [row["line_id"] for row in data["lines"] if not row["quoted"]] == ["L27"]
    assert not any("BRC" in p or "FSSC" in p for p in paragraphs)  # Q1 unanswered -> Partial


def test_northbox_photo_and_golden_boxes():
    data = golden("northbox")
    image = Image.open(reply_path("northbox"))
    assert image.format == "JPEG" and list(image.size) == data["image_size"]
    assert data["skew_deg"] == 14.0
    assert [row["line_id"] for row in data["lines"] if not row["quoted"]] == ["L08", "L26", "L27"]
    for row in data["lines"]:
        if row["quoted"]:
            left, top, right, bottom = row["bbox_px"]
            assert 0 <= left < right <= image.width and 0 <= top < bottom <= image.height
            assert row["normalised_inr_per_1000"] == pytest.approx(row["raw_value"] * 10)
    ambiguous = data["ambiguous_digit"]
    assert ambiguous["readings"] == ["3,935", "3,985"]
    l05 = next(row for row in data["lines"] if row["line_id"] == "L05")
    assert l05["raw_value"] == 3935 and l05["expected_state"] == "needs_review"
    # The blot is really there: the damaged digit's box is much darker than a clean digit's.
    grey = image.convert("L")
    blot = grey.crop(tuple(ambiguous["digit_bbox_px"]))
    assert blot.getextrema()[0] < 80


def test_anand_eml_is_two_lines_with_two_attachments(rfx: RFx):
    data = golden("anand")
    message = email.message_from_bytes(
        reply_path("anand").read_bytes(), policy=email.policy.default
    )
    body = message.get_body(preferencelist=("plain",)).get_content().strip().splitlines()
    assert body == data["body"] and len(body) == 2
    assert (
        body[0] == "₹42/kg for the 5-ply, 38 for the 3-ply, rest same as last year, freight extra."
    )
    names = [part.get_filename() for part in message.iter_attachments()]
    assert names == ["Anand_Questionnaire_RFX-0147.pdf", "Anand_BRCGS_Certificate.pdf"]

    by_id = {line.line_id: line for line in rfx.line_items}
    for row in data["lines"]:
        line = by_id[row["line_id"]]
        if line.spec.ply in (3, 5):
            rate = 42.0 if line.spec.ply == 5 else 38.0
            expected = rate * line.spec.board_weight_kg_per_pc * 1000
            assert row["how"] == "per_kg"
            assert row["normalised_inr_per_1000"] == pytest.approx(expected, abs=0.01)
        elif line.last_year_price_inr is not None:
            assert row["how"] == "carry_forward"
            assert row["normalised_inr_per_1000"] == line.last_year_price_inr
        else:
            assert row["how"] == "carry_forward_unresolved"
            assert row["expected_state"] == "needs_review"
            assert row["normalised_inr_per_1000"] is None


def test_quality_results_and_three_attachments():
    results = {vendor: golden(vendor)["quality_result"] for vendor in VENDORS}
    assert results == {
        "meridian": "Pass",
        "sribalaji": "Pass",
        "kaveri": "Partial",
        "northbox": "Fail",
        "anand": "Pass",
    }
    for vendor in VENDORS:
        answers = golden(vendor)["questionnaire"]
        mandatory = [a["expected"] for a in answers if a["mandatory"]]
        expected = (
            "Fail" if "fail" in mandatory else "Partial" if "unanswered" in mandatory else "Pass"
        )
        assert results[vendor] == expected
    attachments = sorted(p.name for p in (DATA / "attachments").glob("*.pdf"))
    assert attachments == [
        "Anand_BRCGS_Certificate.pdf",
        "Meridian_Spec_Sheet.pdf",
        "SriBalaji_Ink_MSDS.pdf",
    ]
    for name in attachments:
        assert len(pypdf.PdfReader(DATA / "attachments" / name).pages) >= 1


def test_expected_flags_cover_the_briefs_ugly_edges():
    kinds = {
        vendor: {item["kind"] for item in golden(vendor)["expected_review_items"]}
        for vendor in VENDORS
    }
    assert kinds["meridian"] == set()
    assert kinds["sribalaji"] == {"footnote_condition"}
    assert {"unit_ambiguity", "currency_mismatch", "not_quoted"} <= kinds["kaveri"]
    assert {"ocr_disagreement", "not_quoted"} <= kinds["northbox"]
    assert {"carry_forward_unresolved", "missing_freight"} <= kinds["anand"]


def test_source_check_accepts_real_quotes_in_any_shape_and_rejects_invented_ones():
    """The provenance check must survive how models format a table row, and still catch a
    number that is not in the file."""
    from app.agents.extractor import verify_source
    from app.services import parsers

    pdf = parsers.parse("q.pdf", reply_path("sribalaji").read_bytes())
    l04 = next(row for row in golden("sribalaji")["lines"] if row["line_id"] == "L04")
    value = l04["raw_value"]
    shapes = [
        f"4 | 7 Ply export carton 650 x 450 x 450, 180 GSM | 11,000 | pcs | {l04['raw_text']} | 11,91,520.00",
        f"4 7 Ply export carton 650 x 450 x 450, 180 GSM 11,000 pcs {l04['raw_text']}",
        "7 Ply export carton 650 x 450 x 450, 180 GSM",  # description only, number implied
    ]
    for snippet in shapes:
        assert verify_source(pdf, "page 1, row 4", snippet, value) is True, snippet
    assert verify_source(pdf, "page 1, row 4", shapes[0], value + 1000) is False  # wrong number
    assert verify_source(pdf, "page 1", "9 Ply titanium crate 900 x 900 x 900", value) is False

    docx = parsers.parse("k.docx", reply_path("kaveri").read_bytes())
    assert (
        verify_source(docx, "paragraph 7", "at USD 0.65 per box, packed 100 per bundle", 0.65)
        is True
    )
    assert verify_source(docx, "paragraph 7", "at USD 0.56 per box", 0.56) is False
    declined = "We regret that we cannot offer the honeycomb pallet pad 1200x1000x20"
    assert verify_source(docx, "paragraph 13", declined, None) is True
    assert verify_source(docx, "paragraph 13", "We cannot offer the corner post", None) is False

    eml = parsers.parse("a.eml", reply_path("anand").read_bytes())
    assert verify_source(eml, "email body, line 1", "38 for the 3-ply", 38.0) is True
    photo = parsers.parse("n.jpg", reply_path("northbox").read_bytes())
    assert verify_source(photo, "row 5", "anything", 3935.0) is None  # no text layer to check
