"""The HTTP surface, end to end on the golden-derived data (no model calls)."""

from __future__ import annotations

import io

import pypdf
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from tests.helpers import RFX_ID, seed_exact_reads

BUYER = {"X-Persona": "buyer"}
VP = {"X-Persona": "vp"}


@pytest.fixture
def client():
    seed_exact_reads()
    return TestClient(app)


def test_checkpoint_d_grid_endpoint(client):
    grid = client.get(f"/rfx/{RFX_ID}/grid").json()
    cells = [cell for line in grid["lines"] for cell in line["cells"].values()]
    assert len(cells) == 150
    assert grid["counts"] == {"confident": 141, "needs_review": 5, "not_quoted": 4, "overridden": 0}
    for cell in cells:
        if cell["state"] in ("needs_review", "not_quoted"):
            assert cell["inr_per_base"] is None and cell["extended_inr"] is None
    for vendor in grid["vendors"]:
        counted = sum(
            line["cells"][vendor["vendor_id"]]["extended_inr"] or 0 for line in grid["lines"]
        )
        assert vendor["totals"]["extended_inr"] == pytest.approx(counted, abs=0.5)
    assert (
        "RBI" in grid["assumptions"]["fx"] and grid["assumptions"]["base_unit"] == "₹ per 1,000 pcs"
    )


def test_cell_provenance_endpoint(client):
    detail = client.get(f"/rfx/{RFX_ID}/cells/kaveri/L01").json()
    assert detail["cell"]["state"] == "needs_review"
    assert [r["model"] for r in detail["cell"]["readings"]] == [
        "claude-opus-5-5",
        "gemini-3.1-pro-preview",
    ]
    assert detail["cell"]["readings"][0]["source_ref"]["snippet"] == "USD 0.65 per box"
    item = detail["review_items"][0]
    assert item["kind"] == "unit_ambiguity" and len(item["options"]) == 2
    assert [step["op"] for step in item["options"][0]["normalization_trace"]] == [
        "unit_factor",
        "fx",
    ]


def test_personas_gate_actions(client):
    body = {"value": 54000, "reason": "confirmed by phone"}
    assert (
        client.post(f"/rfx/{RFX_ID}/cells/anand/L01/override", json=body, headers=VP).status_code
        == 403
    )
    assert (
        client.post(f"/rfx/{RFX_ID}/cells/anand/L01/override", json=body, headers=BUYER).status_code
        == 200
    )
    assert (
        client.post(
            f"/rfx/{RFX_ID}/cells/anand/L02/override", json={"reason": "x"}, headers=BUYER
        ).status_code
        == 422
    )
    cell = client.get(f"/rfx/{RFX_ID}/cells/anand/L01").json()
    assert cell["cell"]["state"] == "overridden" and cell["history"][0]["action"] == "cell.override"
    assert (
        client.delete(f"/rfx/{RFX_ID}/cells/anand/L01/override", headers=BUYER).status_code == 200
    )


def test_award_is_blocked_until_the_queue_is_clear(client):
    scenario = client.post(
        f"/rfx/{RFX_ID}/scenarios", headers=BUYER,
        json={"name": "Split, quality-cleared", "filters": {"quality_pass_only": True}},
    ).json()  # fmt: skip
    award = client.post(
        f"/rfx/{RFX_ID}/award", json={"scenario_id": scenario["id"]}, headers=BUYER
    ).json()
    award_id = award["award"]["id"]
    assert award["computed"]["approval_blocked"]
    blocked = client.post(f"/award/{award_id}/approve", json={}, headers=VP)
    assert blocked.status_code == 409 and "review item" in blocked.json()["detail"]

    for item in client.get(f"/rfx/{RFX_ID}/review").json():
        option = item["options"][0]["option_id"]
        done = client.post(
            f"/rfx/{RFX_ID}/review/{item['id']}/resolve", json={"option_id": option}, headers=BUYER
        )
        assert done.status_code == 200, done.text
    assert client.get(f"/rfx/{RFX_ID}/grid").json()["review_open"] == 0
    assert client.post(f"/award/{award_id}/approve", json={}, headers=BUYER).status_code == 403
    sent = client.post(f"/award/{award_id}/send-for-approval", json={}, headers=BUYER).json()
    assert sent["award"]["status"] == "sent_for_approval"
    approved = client.post(f"/award/{award_id}/approve", json={"note": "ok"}, headers=VP).json()
    assert approved["award"]["status"] == "approved"
    actions = [event["action"] for event in approved["audit"]]
    assert actions[:2] == ["award.approved", "award.sent_for_approval"]

    pdf = client.get(f"/rfx/{RFX_ID}/award/export.pdf")
    assert pdf.status_code == 200
    text = " ".join(page.extract_text() for page in pypdf.PdfReader(io.BytesIO(pdf.content)).pages)
    assert "Award recommendation" in text and "Rules applied" in text and "Audit log" in text


def test_exports_open_and_say_what_is_excluded(client):
    xlsx = client.get(f"/rfx/{RFX_ID}/export.xlsx")
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert workbook.sheetnames == ["Comparison", "Provenance", "Review items"]
    sheet = workbook["Comparison"]
    values = [str(c.value) for row in sheet.iter_rows() for c in row if c.value is not None]
    assert "not quoted" in values and any(v.startswith("in review") for v in values)
    assert "NOT in any total" in sheet["A3"].value
    assert workbook["Provenance"].max_row > 250
    csv_text = client.get(f"/rfx/{RFX_ID}/export.csv").content.decode("utf-8-sig")
    assert csv_text.count("\n") == 151 and "needs_review" in csv_text


def test_quality_override_needs_the_evaluator(client):
    url = f"/rfx/{RFX_ID}/quality/kaveri/Q1/override"
    body = {"status": "pass", "reason": "certificate received by email"}
    assert client.post(url, json=body, headers=BUYER).status_code == 403
    assert client.post(url, json=body, headers={"X-Persona": "evaluator"}).status_code == 200
    kaveri = next(
        v for v in client.get(f"/rfx/{RFX_ID}/quality").json() if v["vendor_id"] == "kaveri"
    )
    assert kaveri["result"] == "Pass"
    assert (
        next(a for a in kaveri["answers"] if a["question_id"] == "Q1")["effective_status"] == "pass"
    )


def test_changing_the_fx_assumption_moves_only_the_usd_vendor(client):
    before = {
        v["vendor_id"]: v["totals"]["extended_inr"]
        for v in client.get(f"/rfx/{RFX_ID}/grid").json()["vendors"]
    }
    done = client.put(
        f"/rfx/{RFX_ID}/assumptions",
        json={"fx_rate": 85.0, "fx_source": "RBI reference rate"},
        headers=BUYER,
    )
    assert done.status_code == 200
    after = {
        v["vendor_id"]: v["totals"]["extended_inr"]
        for v in client.get(f"/rfx/{RFX_ID}/grid").json()["vendors"]
    }
    assert after["kaveri"] == pytest.approx(before["kaveri"] * 85.0 / 83.42, rel=1e-4)
    assert {k: v for k, v in after.items() if k != "kaveri"} == {
        k: v for k, v in before.items() if k != "kaveri"
    }
