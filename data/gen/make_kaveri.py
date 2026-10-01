"""Kaveri Boards: a .docx letter with the commercials written as prose.

Breakage (brief section 7): quotes in USD; L01 is quoted "per box" with "packed 100 per
bundle" in the same sentence; 29 of 30 lines (the honeycomb pad is declined in a sentence).
No table, no RFx line numbers.
"""

from __future__ import annotations

from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

import catalog as c
from common import REPLIES, normalise_zip

FILENAME = "Kaveri_Quote_Sep2026.docx"

PRINT_WORDS = {None: "", "1 colour": " with one-colour print", "2 colour": " with two-colour print"}


def size(line: c.Line) -> str:
    return "x".join(str(d) for d in line.dims[:3]) if line.dims else ""


def phrase(line: c.Line) -> str:
    """How Kaveri refers to an item in running text."""
    gsm = f" in {line.gsm} GSM" if line.gsm else ""
    printed = PRINT_WORDS[line.print_]
    special = {
        "L01": "the 5-ply RSC 600x400x400 in 150 GSM KLB",
        "L03": f"the die-cut tray {size(line)}",
        "L05": "the 1200x1000 five-ply sheet",
        "L06": "the twelve-cell partition set for the 400x300x300 box",
        "L07": f"the {size(line)} mailer{printed}",
        "L08": "the 50x50x3 edge protector in 600 mm lengths",
        "L12": f"the shelf-ready tray {size(line)}{printed}",
        "L13": "the 390x290 layer pad",
        "L14": "the 1000x800 three-ply sheet",
        "L20": f"the display shipper {size(line)}{printed}",
        "L21": f"the half-slotted {size(line)}",
        "L22": "the six-cell partition set for the 600x400x400 box",
        "L23": f"the telescopic box {size(line)}, base and lid together",
        "L26": "the 40x40x4 edge protector in 1000 mm lengths",
        "L28": "the 80x80x5 corner post in 900 mm lengths",
        "L29": "the die-cut insert pad 380x280",
    }
    if line.line_id in special:
        return special[line.line_id]
    if line.ply == 7:
        return f"the {size(line)} export carton{printed}"
    return f"the {size(line)}{gsm}{printed}"


def price_text(quote: c.Quote, line: c.Line) -> str:
    unit = "per thousand sets" if line.uom == "sets" else "per thousand"
    return f"USD {quote.raw_value:,.2f} {unit}"


def joined(items: list[str]) -> str:
    return "; ".join(items[:-1]) + "; and " + items[-1] if len(items) > 1 else items[0]


def build() -> dict:
    quotes = c.build_quotes()["kaveri"]
    answers = c.QUALITY["kaveri"]["answers"]

    def offer(ids: list[str]) -> str:
        return joined([f"{phrase(c.LINE_BY_ID[i])} at {price_text(quotes[i], c.LINE_BY_ID[i])}" for i in ids])

    l01 = quotes["L01"]
    l01_text = f"USD {l01.raw_value:.2f} per box"
    paragraphs: list[tuple[str, list[str]]] = [
        ("Ref: KB/EXP/2609/26", []),
        ("25 September 2026", []),
        (f"Mr {c.BUYER['contact']}, {c.BUYER['title']}, {c.BUYER['company']}, Chakan, Pune", []),
        (f"Subject: Offer for corrugated packaging against your enquiry {c.RFX_ID}", []),
        ("Dear Mr Thacker,", []),
        (
            "Thank you for including Kaveri Boards in your enquiry. As an export-oriented unit we "
            "invoice in US dollars, so the prices in this letter are in USD on FOB Chennai terms. "
            "Unless we say otherwise each price is per thousand pieces. We have set the offer out "
            "by product family rather than in your table, because several of your items run on "
            "the same line and are best priced together.",
            [],
        ),
        (
            f"Starting with your main shipper, we can supply {phrase(c.LINE_BY_ID['L01'])} at "
            f"{l01_text}, packed 100 per bundle, FOB Chennai. This is the item on which we expect "
            "to be most competitive, as it runs on our new 2.5 metre corrugator.",
            ["L01"],
        ),
        (
            "For the other five-ply cartons we offer " + offer(["L15", "L16", "L17", "L18", "L19", "L21"]) + ".",
            ["L15", "L16", "L17", "L18", "L19", "L21"],
        ),
        (
            "The remaining five-ply items need die-cutting or assembly and are priced accordingly: "
            + offer(["L03", "L20", "L23", "L22", "L05"]) + ".",
            ["L03", "L20", "L23", "L22", "L05"],
        ),
        (
            "In three-ply, the regular slotted boxes are " + offer(["L02", "L09", "L10", "L11", "L30"]) + ".",
            ["L02", "L09", "L10", "L11", "L30"],
        ),
        (
            "The other three-ply items are " + offer(["L07", "L12", "L13", "L29", "L14", "L06"])
            + ". Our sheet and pad prices are keen because we run them as fill-in work between box orders.",
            ["L07", "L12", "L13", "L29", "L14", "L06"],
        ),
        (
            "For the seven-ply export cartons in 180 GSM we quote " + offer(["L04", "L24", "L25"]) + ".",
            ["L04", "L24", "L25"],
        ),
        (
            "On the ancillary items we can offer " + offer(["L08", "L26", "L28"])
            + ". We regret that we cannot offer the honeycomb pallet pad 1200x1000x20, as we do "
            "not manufacture honeycomb board.",
            ["L08", "L26", "L28", "L27"],
        ),
        (
            f"{c.FREIGHT['kaveri']['statement']} {answers['Q6'][1]}. {answers['Q4'][1]}. "
            "This offer holds for 45 days.",
            [],
        ),
        (
            f"On your supplier questions: {answers['Q3'][1]}, and {answers['Q5'][1]}. "
            f"We confirm that {answers['Q2'][1]}.",
            [],
        ),
        ("We look forward to your order and would be glad to send samples of any item.", []),
        ("Yours sincerely,", []),
        ("Arun Subramanian", []),
        ("Head of Exports, Kaveri Boards Private Limited", []),
    ]  # fmt: skip

    doc = Document()
    doc.core_properties.author = "Arun Subramanian"
    doc.core_properties.title = "Offer KB/EXP/2609/26"
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 9, 25, 9, 30, 0)
    style = doc.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = Pt(11)

    header = doc.sections[0].header.paragraphs[0]
    run = header.add_run("KAVERI BOARDS PRIVATE LIMITED")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(0x0B, 0x5C, 0x4A)
    sub = doc.sections[0].header.add_paragraph(
        "100% Export Oriented Unit · Plot 41, SIPCOT Industrial Park, Sriperumbudur 602105 · "
        "exports@kaveriboards.example"
    )
    sub.runs[0].font.size = Pt(8)

    per_line: dict[str, dict] = {}
    locators: dict[str, int] = {}
    for number, (text, line_ids) in enumerate(paragraphs, start=1):
        paragraph = doc.add_paragraph(text)
        if number >= 6 and number <= 16:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if text.startswith("Subject:"):
            paragraph.runs[0].bold = True
        for line_id in line_ids:
            locators[line_id] = number

    for line in c.LINES:
        quote = quotes[line.line_id]
        if not quote.quoted:
            per_line[line.line_id] = {
                "locator": f"paragraph {locators[line.line_id]}",
                "raw_text": "We regret that we cannot offer the honeycomb pallet pad 1200x1000x20",
                "vendor_description": "the honeycomb pallet pad 1200x1000x20",
            }
            continue
        per_line[line.line_id] = {
            "locator": f"paragraph {locators[line.line_id]}",
            "raw_text": l01_text if line.line_id == "L01" else price_text(quote, line),
            "vendor_description": phrase(line),
        }

    path = REPLIES / "kaveri" / FILENAME
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    normalise_zip(path)
    return {
        "files": [FILENAME],
        "lines": per_line,
        "unit_ambiguity": {
            "line_id": "L01",
            "locator": f"paragraph {locators['L01']}",
            "text": f"{l01_text}, packed 100 per bundle",
            "vendor_intent": "per_piece",
        },
        "default_unit_statement": "Unless we say otherwise each price is per thousand pieces.",
        "freight_locator": "paragraph 14",
        "questionnaire_locators": {"Q2": "paragraph 15", "Q3": "paragraph 15", "Q5": "paragraph 15", "Q4": "paragraph 14", "Q6": "paragraph 14"},
    }


if __name__ == "__main__":
    out = build()
    print(out["lines"]["L01"], out["lines"]["L27"], out["lines"]["L06"])
