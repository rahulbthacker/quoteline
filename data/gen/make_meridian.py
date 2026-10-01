"""Meridian Packaging: an .xlsx in the vendor's own layout.

Breakage (brief section 7): ignores the RFx template. No RFx line numbers anywhere, items are
regrouped under Meridian's own sections and SKU codes, header cells are merged, prices sit on
the second sheet, and the column is just called "Rate" with the unit stated in a side note.
"""

from __future__ import annotations

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import catalog as c
from common import REPLIES, normalise_zip

FILENAME = "Meridian_Offer_RFX-0147.xlsx"
SHEET = "Commercial Offer"
UNIT_NOTE = (
    "Note: Rate is in INR per piece (each), delivered to your Chakan plant. "
    "GST @ 18% extra. Value column is for annual quantity."
)

SECTIONS = [
    ("A. REGULAR SLOTTED CARTONS", lambda line: line.style in ("rsc", "hsc") and line.ply in (3, 5)),
    ("B. EXPORT CARTONS (7 PLY)", lambda line: line.ply == 7),
    ("C. DIE CUT ITEMS", lambda line: line.style in ("tray", "display", "mailer", "telescopic") or line.line_id == "L29"),
    ("D. SHEETS, PADS & PARTITIONS", lambda line: line.style in ("sheet", "partition") or line.line_id == "L13"),
    ("E. ANCILLARY ITEMS", lambda line: line.style == "other"),
]  # fmt: skip


def describe(line: c.Line) -> tuple[str, str]:
    """Meridian's own SKU code and wording for an RFx line."""
    size = "".join(f"{d // 10:02d}" for d in line.dims[:3]) if line.dims else ""
    board = f"{line.ply} Ply {line.gsm} GSM" if line.ply else ""
    if line.style == "rsc" and line.ply == 7:
        return f"MP-7X-{size}", f"Export Carton {board} (HD)"
    if line.style == "rsc":
        suffix = " KLB" if "KLB" in line.description else ""
        return f"MP-{line.ply}R-{size}", f"RSC Carton {board}{suffix}"
    if line.style == "hsc":
        return f"MP-{line.ply}H-{size}", f"Half Slotted Carton {board}"
    if line.style == "tray":
        name = "Die Cut SRP Tray" if "shelf-ready" in line.description else "Die Cut Tray"
        return f"MP-DT-{size}", f"{name} {board}"
    if line.style == "display":
        return f"MP-DD-{size}", f"Die Cut Display Shipper {board}"
    if line.style == "mailer":
        return f"MP-DM-{size}", f"Die Cut Mailer Box {board}"
    if line.style == "telescopic":
        return f"MP-TB-{size}", f"Telescopic Box, Base + Lid, {board}"
    if line.style == "pad":
        name = "Die Cut Insert Pad" if "insert" in line.description else "Layer Pad"
        return f"MP-PD-{line.dims[0] // 10}{line.dims[1] // 10}", f"{name} {board}"
    if line.style == "sheet":
        return f"MP-SH-{line.ply}{line.dims[0] // 100}{line.dims[1] // 100}", f"Corrugated Sheet {board}"
    if line.style == "partition":
        cells = "12" if "12 cell" in line.description else "6"
        suits = line.description.split("for ")[1].replace("×", "x").replace(" box", "")
        return f"MP-PT-{line.ply}{cells}", f"Partition Assy {cells} Cell {board} (to suit {suits})"
    names = {
        "L08": ("MP-AB-50503", "Angle Board / Edge Protector 50x50x3"),
        "L26": ("MP-AB-40404", "Angle Board / Edge Protector 40x40x4"),
        "L27": ("MP-HC-121020", "Honeycomb Pallet Pad 20 mm thick"),
        "L28": ("MP-CP-80805", "Corner Post 80x80x5"),
    }
    return names[line.line_id]


def size_cells(line: c.Line) -> tuple[int | str, int | str, int | str]:
    if line.dims is None:
        return ("-", "-", "-")
    if line.style == "other" and line.line_id != "L27":
        return (line.dims[3], "-", "-")  # profile is in the description; only length varies
    dims = list(line.dims[:3]) + ["-"] * (3 - len(line.dims[:3]))
    return tuple(dims)


THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="DDEBF7")
WHITE_BOLD = Font(bold=True, color="FFFFFF")


def build() -> dict:
    quotes = c.build_quotes()["meridian"]
    wb = Workbook()
    wb.properties.creator = "Sneha Kulkarni"
    wb.properties.created = wb.properties.modified = datetime(2026, 9, 22, 10, 0, 0)

    # Sheet 1: cover and supplier information (questionnaire answers in Meridian's own words).
    cover = wb.active
    cover.title = "Cover"
    cover.column_dimensions["A"].width = 30
    cover.column_dimensions["B"].width = 95
    cover.merge_cells("A1:B1")
    cover["A1"] = "MERIDIAN PACKAGING PVT LTD"
    cover["A1"].font = Font(bold=True, size=16, color="1F4E78")
    cover.merge_cells("A2:B2")
    cover["A2"] = "Gat No. 212, Chakan-Talegaon Road, Chakan, Pune 410501 · GSTIN 27AAECM4471K1Z3"
    rows = [
        ("Offer ref", "MP/Q/26-27/0412"),
        ("Date", "22 September 2026"),
        ("To", f"{c.BUYER['contact']}, {c.BUYER['title']}, {c.BUYER['company']}"),
        ("Subject", f"Our offer against your enquiry {c.RFX_ID}: {c.RFX_TITLE}"),
        ("", ""),
        ("", "Dear Sir, thank you for the enquiry. Our item-wise offer is on the next sheet in our standard format. We have quoted for the complete requirement."),
        ("", ""),
        ("SUPPLIER INFORMATION", ""),
    ]  # fmt: skip
    info_labels = {
        "Q1": "Certifications",
        "Q2": "Inks and adhesives",
        "Q3": "Manufacturing locations",
        "Q4": "Lead time",
        "Q5": "Installed capacity",
        "Q6": "Payment terms",
        "Q7": "Recycled content / FSC",
        "Q8": "Safety stock",
    }
    answers = c.QUALITY["meridian"]["answers"]
    questionnaire_locators = {}
    for label, value in rows:
        cover.append([label, value])
    for qid, label in info_labels.items():
        cover.append([label, answers[qid][1]])
        questionnaire_locators[qid] = f"Cover!B{cover.max_row}"
    cover.append(["", ""])
    cover.append(["Freight", c.FREIGHT["meridian"]["statement"]])
    freight_locator = f"Cover!B{cover.max_row}"
    cover.append(["Validity", "This offer is valid for 60 days from the date above."])
    cover.append(["", ""])
    cover.append(["", "For Meridian Packaging Pvt Ltd"])
    cover.append(["", "Sneha Kulkarni, Key Accounts · sales@meridianpackaging.example · +91 20 5550 0142"])
    for row in cover.iter_rows(min_row=3):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    cover["A10"].font = Font(bold=True, color="1F4E78")

    # Sheet 2: the prices.
    ws = wb.create_sheet(SHEET)
    widths = [5, 15, 50, 8, 8, 8, 10, 13, 10, 15, 3, 34]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.merge_cells("A1:J1")
    ws["A1"] = "MERIDIAN PACKAGING PVT LTD · COMMERCIAL OFFER"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
    ws.merge_cells("A2:J2")
    ws["A2"] = f"Ref MP/Q/26-27/0412 dated 22.09.2026 · Customer: {c.BUYER['company']} · Enquiry {c.RFX_ID}"
    ws.merge_cells("A3:J3")
    ws["A3"] = "Items are listed under our product codes. Sizes are internal dimensions."
    ws["A3"].font = Font(italic=True, color="666666")

    # Two-row merged header.
    ws.merge_cells("A5:A6")
    ws.merge_cells("B5:C5")
    ws.merge_cells("D5:F5")
    for col in "GHIJ":
        ws.merge_cells(f"{col}5:{col}6")
    ws["A5"], ws["B5"], ws["D5"] = "Sr", "Item", "Size (mm)"
    ws["G5"], ws["H5"], ws["I5"], ws["J5"] = "Printing", "Annual Qty (nos)", "Rate", "Value (₹)"
    ws["B6"], ws["C6"], ws["D6"], ws["E6"], ws["F6"] = "Code", "Description", "L", "W", "H"
    for row in ws["A5:J6"]:
        for cell in row:
            cell.font, cell.fill, cell.border = WHITE_BOLD, HEAD_FILL, BOX
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # The unit lives in a side note, not in the header.
    ws.merge_cells("L5:L9")
    ws["L5"] = UNIT_NOTE
    ws["L5"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["L5"].font = Font(italic=True, size=9)
    ws["L5"].fill = PatternFill("solid", fgColor="FFF2CC")

    per_line: dict[str, dict] = {}
    row = 7
    serial = 0
    grand_total = 0.0
    for title, belongs in SECTIONS:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        ws.cell(row, 1, title).font = Font(bold=True)
        for col in range(1, 11):
            ws.cell(row, col).fill = SECTION_FILL
            ws.cell(row, col).border = BOX
        row += 1
        members = sorted((line for line in c.LINES if belongs(line)), key=lambda line: describe(line)[0])
        subtotal = 0.0
        for line in members:
            serial += 1
            code, description = describe(line)
            quote = quotes[line.line_id]
            value = round(quote.raw_value * line.qty, 2)
            subtotal += value
            printing = {None: "Plain", "1 colour": "1 Col", "2 colour": "2 Col"}[line.print_]
            cells = [serial, code, description, *size_cells(line), printing, line.qty, quote.raw_value, value]
            for col, cell_value in enumerate(cells, start=1):
                cell = ws.cell(row, col, cell_value)
                cell.border = BOX
                if col in (8, 10):
                    cell.number_format = "#,##0"
                if col == 9:
                    cell.number_format = "0.00"
                if col in (4, 5, 6, 7):
                    cell.alignment = Alignment(horizontal="center")
            per_line[line.line_id] = {
                "locator": f"{SHEET}!I{row}",
                "raw_text": f"{quote.raw_value:.2f}",
                "vendor_description": f"{code} {description}",
            }
            row += 1
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        ws.cell(row, 1, f"Sub-total {title[:1]}").alignment = Alignment(horizontal="right")
        ws.cell(row, 10, round(subtotal, 2)).number_format = "#,##0"
        ws.cell(row, 1).font = ws.cell(row, 10).font = Font(bold=True)
        grand_total += subtotal
        row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    ws.cell(row, 1, "TOTAL ANNUAL VALUE (₹, before GST)").alignment = Alignment(horizontal="right")
    ws.cell(row, 10, round(grand_total, 2)).number_format = "#,##0"
    ws.cell(row, 1).font = ws.cell(row, 10).font = Font(bold=True, size=12)
    ws.freeze_panes = "A7"

    # Sheet 3: terms.
    terms = wb.create_sheet("Terms")
    terms.column_dimensions["A"].width = 110
    for text in [
        "TERMS AND CONDITIONS",
        "1. Prices are firm for the contract period subject to kraft paper index movement within ±5%.",
        "2. Delivery: as per weekly schedule, minimum drop 2 tonnes.",
        "3. Tolerance on quantity ±5%. Tolerance on board GSM ±5%, dimensions ±3 mm.",
        "4. Printing blocks and dies: one-time cost at actuals for new artwork.",
        "5. Payment: 60 days from GRN.",
        "6. This offer is valid for 60 days.",
    ]:
        terms.append([text])
    terms["A1"].font = Font(bold=True)

    path = REPLIES / "meridian" / FILENAME
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    normalise_zip(path)
    return {
        "files": [FILENAME],
        "lines": per_line,
        "unit_note": {"locator": f"{SHEET}!L5", "text": UNIT_NOTE},
        "freight_locator": freight_locator,
        "questionnaire_locators": questionnaire_locators,
    }


if __name__ == "__main__":
    print(build()["lines"]["L01"])
